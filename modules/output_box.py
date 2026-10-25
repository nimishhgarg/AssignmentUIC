
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from core.styles import AppStyles
from .body import ASSUIC_NL, add_paragraph


def _set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _set_cell_margins(cell, v=120):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for tag in ("top", "start", "bottom", "end"):
        node = OxmlElement(f"w:{tag}")
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tcPr.append(mar)


def _remove_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tblPr.append(borders)


def add_output_box(document: Document, lines: list, styles: AppStyles):
    if not lines:
        return
    section = document.sections[0]
    width = section.page_width - section.left_margin - section.right_margin

    table = document.add_table(rows=1, cols=1)
    table.autofit = False
    table.columns[0].width = width
    cell = table.rows[0].cells[0]
    cell.width = width
    _remove_borders(table)
    trPr = table.rows[0]._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))
    _set_cell_background(cell, styles.output.hex("Background-Color", "272727"))
    _set_cell_margins(cell)

    bold, underline = styles.output.bold_underline("Property", "Regular")
    cell.paragraphs[0].text = ""
    for i, line in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(line if line.strip() else "\u00A0")
        run.font.name = styles.output.font("Font", "Consolas")
        run.font.size = styles.output.size("Size")
        run.font.color.rgb = styles.output.color("Text-Color")
        run.bold = bold
        run.underline = underline


def add_assuic_output(document: Document, raw: str, styles: AppStyles):
    lines = [seg.strip() for seg in (raw or "").split(ASSUIC_NL)]
    add_output_box(document, lines, styles)


def add_written_output(document: Document, text: str, styles: AppStyles):
    add_paragraph(document, text, styles)


def add_assuic_written_output(document: Document, raw: str, styles: AppStyles):
    lines = [seg.strip() for seg in (raw or "").split(ASSUIC_NL) if seg.strip()]
    add_written_output(document, "\n\n".join(lines), styles)


def add_image_output(document: Document, image_path: str, styles: AppStyles):
    section = document.sections[0]
    content_width = section.page_width - section.left_margin - section.right_margin
    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_picture(image_path, width=content_width)


def add_output(document: Document, styles: AppStyles, mode: str = "terminal", **kwargs):
    if mode == "terminal":
        add_assuic_output(document, kwargs.get("raw", ""), styles)
    elif mode == "written":
        add_assuic_written_output(document, kwargs.get("raw", ""), styles)
    elif mode == "image":
        add_image_output(document, kwargs["image_path"], styles)
    else:
        raise ValueError(f"unknown output mode: {mode!r}")
