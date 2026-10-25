
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from core.styles import AppStyles






def _remove_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tblPr.append(borders)


def _set_cell_margins(cell, top=0, start=0, bottom=0, end=0):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for tag, val in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = OxmlElement(f"w:{tag}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tcPr.append(mar)


def _set_table_width(table, width_twips):
    tblPr = table._tbl.tblPr
    for existing in tblPr.findall(qn("w:tblW")):
        tblPr.remove(existing)
    tblW = OxmlElement("w:tblW")
    tblW.set(qn("w:type"), "dxa")
    tblW.set(qn("w:w"), str(width_twips))
    tblPr.append(tblW)


def _set_table_indent_zero(table):
    tblPr = table._tbl.tblPr
    for existing in tblPr.findall(qn("w:tblInd")):
        tblPr.remove(existing)
    tblInd = OxmlElement("w:tblInd")
    tblInd.set(qn("w:type"), "dxa")
    tblInd.set(qn("w:w"), "0")
    tblPr.append(tblInd)


def _set_col_widths(table, widths):
    table.autofit = False
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = width
    for i, width in enumerate(widths):
        table.columns[i].width = width
    total_twips = sum(int(w) for w in widths) // 635  
    _set_table_width(table, total_twips)
    _set_table_indent_zero(table)






def add_title(document: Document, title_text: str, styles: AppStyles):
    p = document.add_paragraph()
    p.alignment = styles.title.align("Title-Align", "center")
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    bold, underline = styles.title.bold_underline("Title-Property", "Bold Underline")
    run = p.add_run(title_text)
    run.font.name = styles.title.font("Title-Font")
    run.font.size = styles.title.size("Title-Size")
    run.font.color.rgb = styles.title.color("Title-Color")
    run.bold = bold
    run.underline = underline






def _field_cell(cell, label, value, styles: AppStyles, space_after_pt, align):
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after_pt)

    if label is None:
        return  

    p.paragraph_format.left_indent = Pt(0)
    p.paragraph_format.right_indent = Pt(0)

    label_bold, _ = styles.student.bold_underline("Label-Property", "Bold")
    label_run = p.add_run(f"{label}: ")
    label_run.font.name = styles.student.font("Label-Font")
    label_run.font.size = styles.student.size("Label-Size")
    label_run.font.color.rgb = styles.student.color("Label-Color")
    label_run.bold = label_bold

    value_bold, _ = styles.student.bold_underline("Value-Property", "Regular")
    value_run = p.add_run(str(value) if value is not None else "")
    value_run.font.name = styles.student.font("Value-Font")
    value_run.font.size = styles.student.size("Value-Size")
    value_run.font.color.rgb = styles.student.color("Value-Color")
    value_run.bold = value_bold


def _pair_up(menu: dict, blank_marker: str):
    items = list(menu.items())
    rows = []
    i = 0
    while i < len(items):
        left_label, left_value = items[i]
        if left_value == blank_marker:
            left_label, left_value = None, None
        i += 1

        right_label, right_value = (None, None)
        if i < len(items):
            right_label, right_value = items[i]
            if right_value == blank_marker:
                right_label, right_value = None, None
            i += 1

        rows.append(((left_label, left_value), (right_label, right_value)))
    return rows


def add_student_menu(document: Document, menu: dict, styles: AppStyles):
    blank_marker = styles.student.raw("Blank-Marker", "&assuicblank")
    rows_data = _pair_up(menu, blank_marker)
    if not rows_data:
        return

    section = document.sections[0]
    content_width = section.page_width - section.left_margin - section.right_margin
    half = int(content_width / 2)

    space_after = int(styles.student.raw("Row-Spacing-After", 8))
    R = WD_ALIGN_PARAGRAPH.RIGHT
    L = WD_ALIGN_PARAGRAPH.LEFT

    table = document.add_table(rows=len(rows_data), cols=2)
    _remove_table_borders(table)
    _set_col_widths(table, [half, content_width - half])

    for row_idx, (left, right) in enumerate(rows_data):
        row = table.rows[row_idx]
        for cell in row.cells:
            _set_cell_margins(cell)

        left_blank = left[0] is None
        right_blank = right[0] is None

        if right_blank and not left_blank:
            
            
            
            merged = row.cells[0].merge(row.cells[1])
            _field_cell(merged, left[0], left[1], styles, space_after, L)
        elif left_blank and not right_blank:
            merged = row.cells[0].merge(row.cells[1])
            _field_cell(merged, right[0], right[1], styles, space_after, R)
        else:
            _field_cell(row.cells[0], left[0], left[1], styles, space_after, L)
            _field_cell(row.cells[1], right[0], right[1], styles, space_after, R)






def add_student_menu_lr(document: Document, left: dict, right: dict, styles: AppStyles):
    left_items = list(left.items())
    right_items = list(right.items())
    row_count = max(len(left_items), len(right_items))
    if row_count == 0:
        return

    section = document.sections[0]
    content_width = section.page_width - section.left_margin - section.right_margin
    half = int(content_width / 2)
    space_after = int(styles.student.raw("Row-Spacing-After", 8))
    L, R = WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.RIGHT

    table = document.add_table(rows=row_count, cols=2)
    _remove_table_borders(table)
    _set_col_widths(table, [half, content_width - half])

    for i in range(row_count):
        row = table.rows[i]
        for cell in row.cells:
            _set_cell_margins(cell)
        left_label, left_value = left_items[i] if i < len(left_items) else (None, None)
        right_label, right_value = right_items[i] if i < len(right_items) else (None, None)

        if left_label and not right_label:
            merged = row.cells[0].merge(row.cells[1])
            _field_cell(merged, left_label, left_value, styles, space_after, L)
        elif right_label and not left_label:
            merged = row.cells[0].merge(row.cells[1])
            _field_cell(merged, right_label, right_value, styles, space_after, R)
        else:
            _field_cell(row.cells[0], left_label, left_value, styles, space_after, L)
            _field_cell(row.cells[1], right_label, right_value, styles, space_after, R)


def build_student_section_lr(document: Document, title_text: str, left: dict, right: dict, styles: AppStyles):
    pre_title_space = styles.title.size("Pre-Title-Space", "19pt")
    spacer0 = document.add_paragraph()
    spacer0.paragraph_format.space_before = Pt(0)
    spacer0.paragraph_format.space_after = pre_title_space
    spacer0.add_run("").font.size = Pt(1)

    add_title(document, title_text, styles)

    gap = int(styles.student.raw("Title-To-Menu-Spacing", 5))
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(0)
    spacer.add_run("").font.size = Pt(gap)

    add_student_menu_lr(document, left, right, styles)


def build_student_section(document: Document, title_text: str, menu: dict, styles: AppStyles):
    pre_title_space = styles.title.size("Pre-Title-Space", "19pt")
    spacer0 = document.add_paragraph()
    spacer0.paragraph_format.space_before = Pt(0)
    spacer0.paragraph_format.space_after = pre_title_space
    spacer0.add_run("").font.size = Pt(1)

    add_title(document, title_text, styles)

    gap = int(styles.student.raw("Title-To-Menu-Spacing", 5))
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(0)
    spacer_run = spacer.add_run("")
    spacer_run.font.size = Pt(gap)

    add_student_menu(document, menu, styles)
    
    
    
    
