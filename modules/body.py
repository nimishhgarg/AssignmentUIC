
from docx import Document
from docx.shared import Pt, Inches

from core.styles import AppStyles

ASSUIC_NL = "&assuicnl"
ASSUIC_BULLET = "&assuicbullet"
ASSUIC_INDENT = "&assuicin"


def add_paragraph(document: Document, text: str, styles: AppStyles):
    if not text:
        return
    bold, underline = styles.body.bold_underline("Body-Property", "Regular")
    align = styles.body.align("Body-Align", "justify")
    gap = int(styles.body.raw("Paragraph-Spacing", 8))
    paras = [t.strip() for t in text.split("\n\n") if t.strip()]
    min_lines = int(styles.body.raw("Min-Lines-With-Heading", 3))
    for i, para_text in enumerate(paras):
        p = document.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(gap if i < len(paras) - 1 else 0)
        if i < min_lines - 1 and i < len(paras) - 1:
            p.paragraph_format.keep_with_next = True
        run = p.add_run(para_text.replace("\n", " "))
        run.font.name = styles.body.font("Body-Font")
        run.font.size = styles.body.size("Body-Size")
        run.font.color.rgb = styles.body.color("Body-Color")
        run.bold = bold
        run.underline = underline


def add_bullets(document: Document, items: list, styles: AppStyles, keep_together: bool = False):
    if not items:
        return
    glyph = styles.body.raw("Bullet-Glyph", "\u25AA")
    bold, underline = styles.body.bold_underline("Body-Property", "Regular")
    font_name = styles.body.font("Body-Font")
    font_size = styles.body.size("Body-Size")
    color = styles.body.color("Body-Color")
    indent = styles.body.length("Bullet-Indent", "0.18in")

    for i, item in enumerate(items):
        p = document.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = indent
        pf.first_line_indent = -indent
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        try:
            pf.tab_stops.add_tab_stop(indent)
        except Exception:
            pass
        min_lines = int(styles.body.raw("Min-Lines-With-Heading", 3))
        if (keep_together or i < min_lines - 1) and i < len(items) - 1:
            pf.keep_with_next = True

        bullet_run = p.add_run(f"{glyph}\t")
        bullet_run.font.name = font_name
        bullet_run.font.size = font_size
        bullet_run.font.color.rgb = color
        bullet_run.bold = bold

        text_run = p.add_run(item)
        text_run.font.name = font_name
        text_run.font.size = font_size
        text_run.font.color.rgb = color
        text_run.bold = bold
        text_run.underline = underline


def add_code(document: Document, lines: list, styles: AppStyles):
    if not lines:
        return
    bold, underline = styles.body.bold_underline("Body-Property", "Regular")
    font_name = styles.body.font("Body-Font")
    font_size = styles.body.size("Body-Size")
    color = styles.body.color("Body-Color")
    min_lines = int(styles.body.raw("Min-Lines-With-Heading", 3))
    for i, line in enumerate(lines):
        p = document.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        if i < min_lines - 1 and i < len(lines) - 1:
            p.paragraph_format.keep_with_next = True
        run = p.add_run(line if line.strip() else "\u00A0")
        run.font.name = font_name
        run.font.size = font_size
        run.font.color.rgb = color
        run.bold = bold
        run.underline = underline






def _assuic_lines(raw: str) -> list:
    if not raw:
        return []
    return [seg.strip() for seg in raw.split(ASSUIC_NL)]


def add_assuic_bullets(document: Document, raw: str, styles: AppStyles, keep_together: bool = False):
    items = []
    for line in _assuic_lines(raw):
        if not line:
            continue
        if line.startswith(ASSUIC_BULLET):
            line = line[len(ASSUIC_BULLET):].strip()
        items.append(line)
    add_bullets(document, items, styles, keep_together=keep_together)


def add_assuic_paragraph(document: Document, raw: str, styles: AppStyles):
    paras = [line for line in _assuic_lines(raw) if line]
    add_paragraph(document, "\n\n".join(paras), styles)


def add_assuic_code(document: Document, raw: str, styles: AppStyles):
    lines = []
    for raw_line in (raw or "").split(ASSUIC_NL):
        indent_count = 0
        while raw_line.startswith(ASSUIC_INDENT):
            indent_count += 1
            raw_line = raw_line[len(ASSUIC_INDENT):]
        lines.append(("    " * indent_count) + raw_line)
    add_code(document, lines, styles)
