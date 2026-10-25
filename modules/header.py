
from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from core.styles import AppStyles

RESOURCES_DIR = Path(__file__).resolve().parent.parent / "resources"
LOGO_PRIMARY = RESOURCES_DIR / "logo_primary.jpg"    
LOGO_SECONDARY = RESOURCES_DIR / "logo_secondary.png"  






def _set_page_background(document: Document, hex_color: str):
    if hex_color.upper() in ("FFFFFF", "FFF"):
        return
    bg = OxmlElement("w:background")
    bg.set(qn("w:color"), hex_color)
    document.element.insert(0, bg)


def _anchor_picture(paragraph, image_path, width, height, x_offset_emu, y_offset_emu):
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=width, height=height)

    drawing = run._r.find(qn("w:drawing"))
    inline = drawing.find(qn("wp:inline"))

    anchor = OxmlElement("wp:anchor")
    for attr, val in {
        "distT": "0", "distB": "0", "distL": "0", "distR": "0",
        "simplePos": "0", "relativeHeight": "1", "behindDoc": "0",
        "locked": "0", "layoutInCell": "1", "allowOverlap": "1",
    }.items():
        anchor.set(attr, val)

    simple_pos = OxmlElement("wp:simplePos")
    simple_pos.set("x", "0")
    simple_pos.set("y", "0")
    anchor.append(simple_pos)

    pos_h = OxmlElement("wp:positionH")
    pos_h.set("relativeFrom", "margin")
    off_h = OxmlElement("wp:posOffset")
    off_h.text = str(int(x_offset_emu))
    pos_h.append(off_h)
    anchor.append(pos_h)

    pos_v = OxmlElement("wp:positionV")
    pos_v.set("relativeFrom", "page")
    off_v = OxmlElement("wp:posOffset")
    off_v.text = str(int(y_offset_emu))
    pos_v.append(off_v)
    anchor.append(pos_v)

    wrap_none = OxmlElement("wp:wrapNone")
    anchor.append(wrap_none)

    
    
    for child in list(inline):
        anchor.append(child)

    drawing.replace(inline, anchor)






def setup_page(document: Document, styles: AppStyles):
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)

    margin = styles.page.length("Margin", "0.5in")
    section.bottom_margin = margin
    section.left_margin = margin
    section.right_margin = margin

    header_distance = styles.page.length("Header-Distance", "0.5in")
    section.header_distance = header_distance
    section.footer_distance = styles.page.length("Footer-Distance", "0.5in")

    
    
    
    
    
    
    gap = styles.page.length("Content-Gap-After-Header", "5px")
    tallest_logo = max(
        styles.logo.length("Logo1-Height"),
        styles.logo.length("Logo2-Height"),
    )
    section.top_margin = header_distance + tallest_logo + gap

    _set_page_background(document, styles.page.hex("Background-Color", "FFFFFF"))






def build_header(document: Document, styles: AppStyles):
    section = document.sections[0]
    header = section.header
    header.is_linked_to_previous = False

    p = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)

    content_width = section.page_width - section.left_margin - section.right_margin
    header_distance = section.header_distance

    logo1_w = styles.logo.length("Logo1-Width")
    logo1_h = styles.logo.length("Logo1-Height")
    logo2_w = styles.logo.length("Logo2-Width")
    logo2_h = styles.logo.length("Logo2-Height")
    tallest = max(logo1_h, logo2_h)

    
    
    
    
    
    
    
    
    logo1_y = int(header_distance) + int(tallest) - int(logo1_h)
    logo2_y = int(header_distance) + int(tallest) - int(logo2_h)
    logo2_x = int(content_width) - int(logo2_w)

    _anchor_picture(p, LOGO_PRIMARY, logo1_w, logo1_h, x_offset_emu=0, y_offset_emu=logo1_y)
    _anchor_picture(p, LOGO_SECONDARY, logo2_w, logo2_h, x_offset_emu=logo2_x, y_offset_emu=logo2_y)
