
from docx import Document
from docx.shared import Pt

from core.styles import AppStyles


def add_heading(document: Document, text: str, styles: AppStyles, is_first: bool = False):
    key = "First-Heading-Space" if is_first else "Pre-Heading-Space"
    default = "15pt" if is_first else "5pt"
    gap = styles.heading.size(key, default)
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = gap
    spacer.add_run("").font.size = Pt(1)

    post_gap = styles.heading.size("Post-Heading-Space", "1px")

    p = document.add_paragraph()
    p.alignment = styles.heading.align("Heading-Align", "start")
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = post_gap  
    p.paragraph_format.keep_with_next = True  
    bold, underline = styles.heading.bold_underline("Heading-Property", "Bold")
    run = p.add_run(text)
    run.font.name = styles.heading.font("Heading-Font")
    run.font.size = styles.heading.size("Heading-Size")
    run.font.color.rgb = styles.heading.color("Heading-Color")
    run.bold = bold
    run.underline = underline
